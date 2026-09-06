"""Regression checks on rendered ink and editable Office packages.

Run after build.py. An optional PDF argument lets the same check demonstrate
the original circle-alignment failure without rebuilding it.
"""
from pathlib import Path
import json
import sys
import zipfile
import pymupdf
from PIL import Image
from lxml import etree
from pptx import Presentation
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parent/'dist'
if not __debug__:
    raise RuntimeError('Run layout validation without -O; assertions must remain enabled.')
doc=pymupdf.open(sys.argv[1] if len(sys.argv)>1 else OUT/'Dayova-Reference.pdf')
assert len(doc)==20
checks=[]
regions=[(1,x,y,70,70,'FFFFFF' if i==0 else '1A1A1A') for i,(x,y) in enumerate([(889,477),(1070,381),(906,268),(1070,181)])]
regions += [(8,56+i*300,337,64,64,'1A1A1A') for i in range(4)]
regions += [(16,152+i*400,314,176,176,'5F6B7C') for i in range(3)]
regions += [(18,114,512,314,64,'FFFFFF')]
for page,x,y,w,h,color in regions:
    pix=doc[page-1].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(x,y,x+w,y+h))
    im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    rgb=tuple(bytes.fromhex(color))
    points=[(xx,yy) for yy in range(im.height) for xx in range(im.width)
            if max(abs(a-b) for a,b in zip(im.getpixel((xx,yy)),rgb))<20]
    assert points, f'Label missing on page {page}'
    dy=(min(p[1] for p in points)+max(p[1] for p in points)+1)/6-h/2
    dx=(min(p[0] for p in points)+max(p[0] for p in points)+1)/6-w/2
    # PDF units deliberately equal pixels on the 1280 x 720 source canvas.
    # Raster measurements are divided by the 3x rendering scale above.
    checks.append(dict(page=page,x=x,verticalOffsetCanvasPx=round(dy,3),horizontalOffsetCanvasPx=round(dx,3)))
assert max(abs(c['verticalOffsetCanvasPx']) for c in checks)<=1,checks
assert max(abs(c['horizontalOffsetCanvasPx']) for c in checks)<=2,checks
bullet_checks=[]
for i in range(3):
    for x,w in [(111,530),(660,560)]:
        pix=doc[19].get_pixmap(matrix=pymupdf.Matrix(3,3),clip=pymupdf.Rect(x,315+i*99,x+w,374+i*99))
        im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
        ys=[y for y in range(im.height) for xx in range(im.width) if min(im.getpixel((xx,y)))>235]
        assert ys, 'Missing dark-content label'
        dy=(min(ys)+max(ys)+1)/6-29.5
        assert abs(dy)<=1,dy
        bullet_checks.append(dict(row=i+1,x=x,verticalOffsetCanvasPx=round(dy,3)))
for pg in doc:
    for b in pg.get_text('dict')['blocks']:
        for line in b.get('lines',[]):
            for span in line['spans']:
                # Allow font descender bounds at the footer; visible ink is inside.
                assert span['bbox'][0]>=0 and span['bbox'][2]<=1281
                assert span['bbox'][1]>=0 and span['bbox'][3]<=721
for filename in ['Dayova-Presentation.pptx','Dayova-Presentation.potx','Dayova-Workbook.xlsx']:
    with zipfile.ZipFile(OUT/filename) as z:
        assert z.testzip() is None
        for n in z.namelist():
            if n.endswith('.xml'):etree.fromstring(z.read(n))
prs=Presentation(OUT/'Dayova-Presentation.pptx')
assert len(prs.slides)==20
for slide in prs.slides:
    assert slide.notes_slide.notes_text_frame.text.strip()
    for sh in slide.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=prs.slide_width+10 and sh.top+sh.height<=prs.slide_height+10
for i in [0,7,15,17]:
    assert any(s.has_text_frame and s.text_frame.vertical_anchor==3 for s in prs.slides[i].shapes)
wb=load_workbook(OUT/'Dayova-Workbook.xlsx')
assert wb.sheetnames==['Start','Dashboard','Lernschritte','Budget','Marke']
for cell in ['B10','E10','H10','B22','E22','H22']:
    assert wb['Dashboard'][cell].alignment.horizontal=='left'
for cell in ['G7','G106','I7','I106']:
    assert wb['Lernschritte'][cell].alignment.horizontal=='center'
result=dict(pdfPages=20,officePackages=3,opticalLabels=checks,bulletRows=bullet_checks,nativePowerPointRender='blocked by window ownership error; not covered by these tests')
(OUT/'layout-check.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8',newline='\n')
doc[7].get_pixmap().save(OUT/'Dayova-Process.png')
print(json.dumps(result,indent=2))
