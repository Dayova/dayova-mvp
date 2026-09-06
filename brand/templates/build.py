"""Build the editable Dayova Office / Canva kit. Run from any directory."""
from pathlib import Path
import base64, copy, html, json, zipfile
from datetime import date, timedelta
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageFont
from functools import lru_cache

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent / 'dist'
OUT.mkdir(exist_ok=True)
C = dict(bg='F6F6F4', ink='1A1A1A', muted='5F6B7C', appMuted='697586', cyan='00BAFF', strong='00A0E6', ice='F1F7FB', white='FFFFFF', border='DCE6EE', purple='5856D6', lavender='EEECFF', green='34C759', palegreen='EAFFF1', orange='FF9500', paleorange='FFECD6', dark='212325')
W,H=1280,720
pages=[]
def page(label, title=None, subtitle=None, dark=False, note=''):
    p={'label':label,'bg':C['dark'] if dark else C['bg'],'elements':[], 'note':note}
    pages.append(p)
    text(p,'dayova.',56,32,200,42,28,'white' if dark else 'ink',True)
    text(p,label.upper(),900,46,324,24,12,'white' if dark else 'muted',False,'right')
    if title: text(p,title,56,126,1168,100,44,'white' if dark else 'ink',True)
    if subtitle: text(p,subtitle,56,229,1120,52,20,'white' if dark else 'muted')
    rect(p,56,672,1168,1,'border')
    text(p,'DAYOVA  /  VORLAGEN · 2026',56,686,800,20,10,'white' if dark else 'muted')
    text(p,f'{len(pages):02}',1164,682,60,25,13,'white' if dark else 'muted',False,'right')
    return p
def text(p,s,x,y,w,h,size=22,color='ink',bold=False,align='left'):
    p['elements'].append(dict(kind='text',text=s,x=x,y=y,w=w,h=h,size=size,color=C.get(color,color),bold=bold,align=align))

@lru_cache
def ink_bounds(s, bold):
    """Visible glyph bounds relative to the baseline, in em units."""
    font=ImageFont.truetype(str(ROOT/'assets/fonts'/('Poppins-SemiBold.ttf' if bold else 'Poppins-Regular.ttf')),1000)
    left,top,right,bottom=font.getbbox(s,anchor='ls')
    return top/1000,bottom/1000

def centered_text(p,s,x,y,w,h,size=22,color='ink',bold=True,align='center'):
    """Center visible ink, not the font's asymmetric ascender/descender box."""
    text(p,s,x,y,w,h,size,color,bold,align)
    p['elements'][-1]['optical_center']=True
def rect(p,x,y,w,h,color='white',radius=0):
    p['elements'].append(dict(kind='rect',x=x,y=y,w=w,h=h,color=C.get(color,color),radius=radius))
def circle(p,x,y,d,color='cyan'):
    p['elements'].append(dict(kind='circle',x=x,y=y,w=d,h=d,color=C.get(color,color)))
def card(p,x,y,w,h,num,title,body,color='white'):
    rect(p,x,y,w,h,color,24);text(p,num,x+28,y+24,w-56,52,32,'ink',True)
    text(p,title,x+28,y+100,w-56,70,25,'ink',True);text(p,body,x+28,y+185,w-56,h-195,19,'muted')

p=page('Titel',note='Titel durch maximal 7 Wörter ersetzen. Untertitel: ein Satz. Absender, Anlass und Datum ergänzen. Die abstrakte Route ist ein Kommunikationsmotiv, keine nachgebaute App-Oberfläche.')
rect(p,816,112,408,524,'cyan',32)
text(p,'Einfach\nloslernen.',56,181,735,195,76,'ink',True)
text(p,'Der Plan steht schon.',60,412,700,58,34,'ink',True)
text(p,'[Anlass / Präsentation]\n[Name] · [Datum]',60,537,660,70,20,'muted')
route=[(889,477),(1070,381),(906,268),(1070,181)]
for (x1,y1),(x2,y2) in zip(route,route[1:]):
    mid=(y1+y2)/2+35
    rect(p,x1+33,mid,4,y1+35-mid,'ink',2)
    rect(p,min(x1,x2)+35,mid-2,abs(x2-x1),4,'ink',2)
    rect(p,x2+33,y2+35,4,mid-y2-35,'ink',2)
for i,(x,y) in enumerate(route):
    circle(p,x,y,70,'ink' if i==0 else 'white');centered_text(p,f'{i+1:02}',x,y,70,70,25,'white' if i==0 else 'ink')
text(p,'SCHRITT FÜR SCHRITT',853,573,330,28,14,'ink',True)

p=page('Anleitung','Eine Vorlage. Viele Möglichkeiten.','Kopieren, Inhalte ersetzen, loslegen.',note='Diese Hilfsfolie vor dem Präsentieren entfernen. Schrift Poppins installieren. Alle Text- und Formelemente sind bearbeitbar. PowerPoint: Folie duplizieren; Canva: Seite duplizieren. Zahlen sind fiktiv.')
for x,n,t,b in [(56,'01','Layout wählen','Wähle die Folie passend zur Aussage. Eine Kernbotschaft pro Seite.'),(456,'02','Inhalte einsetzen','Ersetze eckige Klammern. Prüfe Quellen, Zahlen und Bildrechte.'),(856,'03','Kurz prüfen','Lies in Präsentationsgröße. Prüfe Umbrüche und exportiere ein PDF.')]:card(p,x,310,368,305,n,t,b)

p=page('Agenda','Was heute zählt.','[Ziel des Gesprächs in einem Satz]',note='Agenda: maximal vier Punkte. Ersetze Inhalte, behalte kurze Titel und eindeutige Reihenfolge.')
for i,(t,b) in enumerate([('Ausgangspunkt','Was wir verstehen müssen'),('Unser Ansatz','Wie wir den nächsten Schritt erleichtern'),('Einblick','Was wir zeigen und lernen'),('Nächster Schritt','Was wir gemeinsam entscheiden')]):
    y=306+i*80;text(p,f'0{i+1}',60,y,100,56,32,'ink',True);text(p,t,170,y,430,48,26,'ink',True);text(p,b,650,y+5,555,40,20,'muted')

p=page('Kapitel',dark=True,note='Kapiteltrenner: große Ziffer, kurze Kapitelüberschrift. Keine zusätzliche Liste.')
text(p,'01',56,152,320,205,150,'cyan',True);text(p,'Der nächste\nsinnvolle Schritt.',424,214,800,180,58,'white',True);text(p,'[Kapitel / Leitfrage]',430,483,730,55,23,'white')

p=page('Kernaussage','Weniger planen.\nLeichter anfangen.',note='Kernaussage mit drei Belegen. Die Beispieltexte beschreiben das Produktprinzip, keine nachgewiesenen Wirkungsversprechen.')
rect(p,56,361,1168,256,'white',28)
for x,n,t,b in [(84,'01','Orientierung','Was ist heute dran?'),(474,'02','Ein klarer Start','Wie beginne ich?'),(864,'03','Passend weiter','Was brauche ich danach?')]:
    text(p,n,x,390,290,42,24,'ink',True);text(p,t,x,457,300,40,23,'ink',True);text(p,b,x,518,300,60,18,'muted')

p=page('Zwei Spalten','Von der Frage zum nächsten Schritt.','[Kontext für den Vergleich]',note='Zwei Perspektiven oder Problem / Ansatz. Keine unbelegten absoluten Versprechen.')
for x,t,b,color in [(56,'Die Ausgangslage','[Beschreibe die Situation.]\n\n[Was macht den Start schwer?]\n\n[Was soll einfacher werden?]','white'),(652,'Unser Ansatz','[Beschreibe die Antwort.]\n\n[Wie wird der nächste Schritt klar?]\n\n[Woran erkennen wir Fortschritt?]','ice')]:
    rect(p,x,315,572,313,color,28);text(p,t,x+32,342,505,55,28,'ink',True);text(p,b,x+32,416,505,192,20,'muted')

p=page('Drei Bausteine','Ein Lernweg, der zusammenpasst.','[Drei Gedanken, die die Kernaussage tragen]',note='Drei gleichwertige Bausteine. Kurze Überschriften, je maximal 25 Wörter Erklärung.')
for x,n,t,b,col in [(56,'01','Prüfung','Termin, Themen und verfügbare Zeit geben den Rahmen.','white'),(456,'02','Lernschritt','Ein konkreter Fokus erleichtert den Anfang.','ice'),(856,'03','Rückmeldung','Antworten zeigen, was als Nächstes sinnvoll ist.','lavender')]:card(p,x,313,368,310,n,t,b,col)

p=page('Prozess','Vom Stoff zum Lernschritt.','Ein Ablauf, den man auf einen Blick versteht.',note='Prozessgrafik als bearbeitbare Formen. Kein Screenshot und keine Behauptung über eine konkrete App-Ansicht.')
for i,(t,b) in enumerate([('Eintragen','Prüfung, Themen\nund freie Zeit'),('Plan ansehen','Die nächsten\nLerntage im Blick'),('Starten','Ein klarer Fokus\nfür heute'),('Weiterlernen','Antworten geben\ndie Richtung vor')]):
    x=56+i*300;circle(p,x,337,64);centered_text(p,str(i+1),x,337,64,64,28)
    if i<3:rect(p,x+84,367.5,196,3,'border')
    text(p,t,x,433,268,58,26,'ink',True);text(p,b,x,511,264,90,21,'muted')

p=page('Produkteinblick','Ein klarer Fokus\nfür heute.',note='Referenz: Repository docs/evidence/day-292-onboarding/android-light-intro-daily-guidance-shared.png. Echter aufgezeichneter App-Zustand, nicht frisch live aufgenommen. Vor externer Nutzung auf Aktualität prüfen. Bild mit Quelle ersetzen; Seitenverhältnis beibehalten.')
text(p,'[Erkläre den gezeigten Moment\nin höchstens zwei Sätzen.]',56,345,690,110,25,'muted')
rect(p,56,514,598,100,'ice',22);text(p,'BILDQUELLE IM FOLIENKOMMENTAR\nVor Veröffentlichung aktualisieren.',82,537,548,70,16,'ink')
asset=ROOT/'docs/evidence/day-292-onboarding/android-light-intro-daily-guidance-shared.png'
p['elements'].append(dict(kind='image',path=str(asset),x=894,y=109,w=238,h=531))

p=page('Bild + Text','Ein Bild. Eine klare Aussage.','[Was zeigt das Bild und warum ist es relevant?]',note='Die graue Fläche ist ein Bildplatzhalter. In PowerPoint/Canva eigenes freigegebenes Bild einfügen, proportional beschneiden. Alternativtext und Quelle ergänzen.')
rect(p,56,311,744,318,'ice',28);text(p,'[Bild einsetzen]',130,421,600,64,32,'muted',True,'center');text(p,'FORMAT  2,3 : 1',210,509,440,32,14,'muted',False,'center')
text(p,'[Kernaussage]',844,326,380,70,29,'ink',True);text(p,'[Kurze Einordnung]\n\n[Wichtigstes Detail]\n\n[Quelle / Stand]',844,425,380,195,20,'muted')

p=page('Kennzahlen','Fortschritt sichtbar machen.','Fiktives Beispiel · Werte vor Verwendung ersetzen',note='Fiktive Beispiele. KPI immer mit Definition, Zeitraum und Quelle angeben. Keine realen Dayova Leistungsdaten.')
for x,n,t,b,col in [(56,'24','Lernschritte','im betrachteten Zeitraum','white'),(456,'75 %','Abgeschlossen','18 von 24 Lernschritten','ice'),(856,'30 Min.','Pro Lernschritt','geplante Dauer im Beispiel','white')]:
    rect(p,x,328,368,282,col,28);text(p,n,x+28,361,312,95,61,'ink',True);text(p,t,x+28,487,312,55,23,'ink',True);text(p,b,x+28,558,312,39,15,'muted')

p=page('Diagramm','Eine Entwicklung, die lesbar bleibt.','Abgeschlossene Lernschritte je Woche · Fiktives Beispiel',note='PowerPoint: Rechtsklick auf Diagramm > Daten bearbeiten. Canva: Balken und Werte sind einzelne bearbeitbare Elemente, kein verknüpftes Diagramm. Achse beginnt bei null. Werte 6, 9, 12, 15; Einheit Lernschritte.')
rect(p,56,310,780,329,'white',24)
p['elements'].append(dict(kind='chart',x=86,y=334,w=704,h=263,values=[6,9,12,15],labels=['Woche 1','Woche 2','Woche 3','Woche 4']))
text(p,'[Was fällt auf?]',882,338,340,70,28,'ink',True);text(p,'[Beobachtung]\n\n[Einordnung]\n\n[Nächster Schritt]',882,434,340,172,21,'muted')

p=page('Tabelle','Details, die die Entscheidung tragen.','Fiktives Beispiel · Status immer auch als Text zeigen',note='PowerPoint-Tabelle ist nativ bearbeitbar. Canva-Zellen sind bearbeitbare Texte und Formen. Maximal fünf Zeilen pro Folie; längere Tabellen aufteilen.')
p['elements'].append(dict(kind='table',x=56,y=318,w=1168,h=282,rows=[['Bereich','Beobachtung','Status','Nächster Schritt'],['Planung','[Kurze Beobachtung]','Offen','[Aktion]'],['Lernschritte','[Kurze Beobachtung]','In Arbeit','[Aktion]'],['Rückmeldung','[Kurze Beobachtung]','Erledigt','[Aktion]']]))

p=page('Zeitplan','Ein Plan mit klaren Etappen.','[Zeitraum und Ziel] · Beispielablauf, keine Zusage',note='Phasen mit Verantwortlichen und Termin ersetzen. Keine automatische Projektverfolgung; Linear bleibt das kanonische Issuesystem.')
for i,(t,b) in enumerate([('Vorbereiten','Ziel & Rahmen\n[Name · Datum]'),('Starten','Einführung\n[Name · Datum]'),('Beobachten','Rückmeldung\n[Name · Datum]'),('Entscheiden','Auswertung\n[Name · Datum]')]):
    x=56+i*300;rect(p,x,332,268,16,'cyan' if i==0 else 'border',8);text(p,f'0{i+1}',x,381,260,60,35,'ink',True);text(p,t,x,469,268,55,25,'ink',True);text(p,b,x,542,268,84,19,'muted')

p=page('Zitat',dark=True,note='Platzhalter, kein echtes Testimonial. Nur freigegebene Originalzitate verwenden, mit Name/Rolle/Datum und dokumentierter Freigabe. Platzhalter komplett ersetzen.')
text(p,'“',50,97,180,180,140,'cyan',True);text(p,'[Ein kurzer Gedanke,\nder in Erinnerung\nbleiben soll.]',210,202,930,248,52,'white',True);text(p,'[Name]  ·  [Rolle / Kontext]\n[Quelle · Datum · Freigabe]',218,531,850,84,19,'white')

p=page('Menschen','Die Menschen hinter dem nächsten Schritt.','[Team / Ansprechpartner]',note='Personen sind Platzhalter. Freigegebene Fotos ersetzen die Initialen. Keine erfundenen Namen oder Rollen.')
for x in [56,456,856]:
    circle(p,x+96,314,176,'ice');centered_text(p,'[AB]',x+96,314,176,176,40,'muted');text(p,'[Name]',x+25,518,318,49,29,'ink',True,'center');text(p,'[Rolle / Schwerpunkt]',x+12,583,344,40,18,'muted',False,'center')

p=page('Entscheidung','Die Optionen im Überblick.','[Welche Entscheidung steht an?]',note='Kriterien parallel halten. Empfohlene Option mit Begründung hervorheben. Platzhalter ersetzen; keine Preisangaben ohne aktuelle Quelle.')
for x,t,col in [(56,'Option A','white'),(652,'Option B','ice')]:
    rect(p,x,313,572,310,col,28);text(p,t,x+30,338,510,50,30,'ink',True);text(p,'[Nutzen]\n\n[Aufwand / Voraussetzung]\n\n[Wichtigste Einschränkung]',x+30,417,510,195,21,'muted')

p=page('Abschluss',note='Abschluss: eine konkrete Handlung, Person und Termin. Kontakt / Link vor Veröffentlichung prüfen.')
rect(p,56,126,1168,500,'cyan',32);text(p,'Was ist unser\nnächster Schritt?',96,166,1080,190,64,'ink',True);text(p,'[Konkrete Handlung]  ·  [Name]  ·  [Termin]',99,401,1050,65,25,'ink');rect(p,99,512,344,64,'ink',32);centered_text(p,'[Aktion benennen]',114,512,314,64,22,'white');text(p,'dayova.com',913,546,261,36,19,'ink',False,'right')

p=page('Markenbausteine','Dayova, konsistent eingesetzt.','Poppins · klare Hierarchie · warmes Weiß · Cyan als Akzent',note='Hilfsfolie vor dem Präsentieren entfernen. Primärfarbe #00BAFF mit #1A1A1A Text. Poppins Regular für Fließtext; SemiBold für Überschriften. Kleine farbige Texte vermeiden. Diese Kommunikationsskala 64/44/28/20 px ändert keine App-Tokens.')
for i,(label,col) in enumerate([('Cyan','cyan'),('Text','ink'),('Fläche','white'),('Hintergrund','bg'),('Theorie','purple')]):
    x=56+i*236;rect(p,x,324,220,119,col,20);text(p,label,x,461,220,35,19,'ink',True);text(p,'#'+C[col],x,505,220,32,17,'muted')
text(p,'64  /  44  /  28  /  20',56,573,620,60,32,'ink',True);text(p,'Titel / Aussage / Abschnitt / Text',726,581,498,60,20,'muted')

p=page('Dunkle Inhaltsfolie','Fokus auf das Wesentliche.','[Eine klare Aussage für den dunklen Modus]',dark=True,note='Dunkle Inhaltsvariante. Weiß auf #212325; Cyan nur als Akzent. Keine Übertragung auf die App-Themesteuerung.')
for i,(t,b) in enumerate([('[Gedanke 1]','[Kurze Erläuterung]'),('[Gedanke 2]','[Kurze Erläuterung]'),('[Gedanke 3]','[Kurze Erläuterung]')]):
    y=332+i*99;circle(p,58,y+4,17);centered_text(p,t,111,y-17,530,59,28,'white',True,'left');centered_text(p,b,660,y-17,560,59,22,'white',False,'left')

prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='Dayova — Präsentationsvorlage';prs.core_properties.author='Dayova';prs.core_properties.subject='20 reusable layouts; fictional examples explicitly labelled'
def rgb(s):return RGBColor.from_string(s)
def box(s,e):return tuple(Inches(e[k]/96) for k in ['x','y','w','h'])
def ppt_text(s,e):
    sh=s.shapes.add_textbox(*box(s,e));tf=sh.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    for i,line in enumerate(e['text'].split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Poppins SemiBold' if e['bold'] else 'Poppins';p.font.size=Pt(e['size']*.75);p.font.color.rgb=rgb(e['color']);p.font.bold=False;p.space_before=Pt(0);p.space_after=Pt(0);p.line_spacing=1.25;p.alignment={'left':PP_ALIGN.LEFT,'right':PP_ALIGN.RIGHT,'center':PP_ALIGN.CENTER}[e['align']]
    if e.get('optical_center'):
        tf.vertical_anchor=MSO_ANCHOR.MIDDLE
        # Poppins ascent/descent are 1.05/-0.35 em. Compensate its
        # typographic midpoint to center the actual label's visible ink.
        top,bottom=ink_bounds(e['text'],e['bold'])
        shift=(-.35-(top+bottom)/2)*e['size']
        tf.margin_top=Inches(max(0,2*shift)/96)
        tf.margin_bottom=Inches(max(0,-2*shift)/96)
        p.line_spacing=1.0
    return sh
for pg in pages:
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(pg['bg']);s.notes_slide.notes_text_frame.text=pg['note']
    for e in pg['elements']:
        kind=e['kind']
        if kind=='text':ppt_text(s,e)
        elif kind in ['rect','circle']:
            sh=s.shapes.add_shape(MSO_SHAPE.OVAL if kind=='circle' else MSO_SHAPE.ROUNDED_RECTANGLE if e.get('radius') else MSO_SHAPE.RECTANGLE,*box(s,e));sh.fill.solid();sh.fill.fore_color.rgb=rgb(e['color']);sh.line.fill.background()
            if kind=='rect' and e.get('radius'):sh.adjustments[0]=min(.5,e['radius']/min(e['w'],e['h']))
        elif kind=='image':s.shapes.add_picture(e['path'],*box(s,e))
        elif kind=='chart':
            d=CategoryChartData();d.categories=e['labels'];d.add_series('Lernschritte',e['values']);ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,*box(s,e),d).chart;ch.has_legend=False;ch.value_axis.minimum_scale=0;ch.value_axis.major_unit=5
            ch.category_axis.tick_labels.font.name='Poppins';ch.category_axis.tick_labels.font.size=Pt(12);ch.value_axis.tick_labels.font.size=Pt(10)
            ch.series[0].format.fill.solid();ch.series[0].format.fill.fore_color.rgb=rgb(C['cyan']);ch.series[0].format.line.fill.background();ch.plots[0].has_data_labels=True;ch.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END;ch.plots[0].data_labels.font.name='Poppins';ch.plots[0].data_labels.font.size=Pt(13)
        elif kind=='table':
            tb=s.shapes.add_table(len(e['rows']),4,*box(s,e)).table
            for r,row in enumerate(e['rows']):
                for c,value in enumerate(row):
                    from pptx.enum.text import MSO_ANCHOR
                    cell=tb.cell(r,c);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=rgb(C['ink'] if r==0 else C['white'] if r%2 else C['ice']);cell.margin_left=Inches(.18);cell.margin_top=cell.margin_bottom=0;cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                    for pp in cell.text_frame.paragraphs:pp.font.name='Poppins';pp.font.size=Pt(14);pp.font.color.rgb=rgb(C['white'] if r==0 else C['ink'])
    for sh in s.shapes:
        if sh.has_text_frame:sh.name=(sh.text[:60] or 'Text')
prs.save(OUT/'Dayova-Presentation.pptx')
# Apply brand theme to Office UI, including charts created later.
with zipfile.ZipFile(OUT/'Dayova-Presentation.pptx') as z:parts={n:z.read(n) for n in z.namelist()}
from lxml import etree
theme=etree.fromstring(parts['ppt/theme/theme1.xml']);ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
for node in theme.findall('.//a:majorFont/a:latin',ns)+theme.findall('.//a:minorFont/a:latin',ns):node.set('typeface','Poppins')
for i,key in enumerate(['cyan','purple','green','orange','strong','muted'],1):
    node=theme.find(f'.//a:clrScheme/a:accent{i}',ns);node.clear();etree.SubElement(node,'{'+ns['a']+'}srgbClr',val=C[key])
parts['ppt/theme/theme1.xml']=etree.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
for name,template in [('Dayova-Presentation.pptx',False),('Dayova-Presentation.potx',True)]:
    with zipfile.ZipFile(OUT/name,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in parts.items():
            if template and n=='[Content_Types].xml':data=data.replace(b'presentationml.presentation.main+xml',b'presentationml.template.main+xml')
            z.writestr(n,data)

def chart_elements(e):
    out=[];q={'elements':out};x,y,w,h=[e[k] for k in ['x','y','w','h']]
    for v in [0,5,10,15,20]:
        yy=y+h-34-v/20*(h-66);rect(q,x+34,yy,w-40,1,'border');text(q,str(v),x,yy-10,27,22,12,'muted')
    for i,(v,label) in enumerate(zip(e['values'],e['labels'])):
        bx=x+70+i*(w-60)/4;bh=v/20*(h-66);rect(q,bx,y+h-34-bh,72,bh,'cyan');text(q,str(v),bx,y+h-64-bh,72,24,16,'ink',True,'center');text(q,label,bx-19,y+h-23,110,24,13,'muted',False,'center')
    return out
def table_elements(e):
    out=[];q={'elements':out};cw=e['w']/4;rh=e['h']/len(e['rows'])
    for r,row in enumerate(e['rows']):
        for c,v in enumerate(row):
            x=e['x']+c*cw;y=e['y']+r*rh;rect(q,x,y,cw,rh,'ink' if r==0 else 'white' if r%2 else 'ice');centered_text(q,v,x+17,y,cw-34,rh,18,'white' if r==0 else 'ink',r==0,'left');out[-1]['center_reference']='Ag'
    return out
def flattened(pg):
    out=[]
    for e in pg['elements']:out.extend(chart_elements(e) if e['kind']=='chart' else table_elements(e) if e['kind']=='table' else [e])
    return out
fontcss=''
for weight,file in [(400,'Poppins-Regular.ttf'),(600,'Poppins-SemiBold.ttf')]:
    blob=base64.b64encode((ROOT/'assets/fonts'/file).read_bytes()).decode();fontcss+=f'@font-face{{font-family:Poppins;font-weight:{weight};src:url(data:font/ttf;base64,{blob})}}'
sections=[]
for pg in pages:
    els=[]
    for e in flattened(pg):
        e=e.copy()
        if e['kind']=='text':e['text']=e['text'].replace('\n\n','\n\u00a0\n')
        if e.get('optical_center'):
            top,bottom=ink_bounds(e.get('center_reference',e['text']),e['bold'])
            # Browser line-height 1.25 gives a Poppins baseline at .975 em.
            e['y']+=e['h']/2-(top+bottom)*e['size']/2-.975*e['size']
            e['h']=e['size']*1.4
        style=f"position:absolute;left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;"
        if e['kind']=='text':els.append(f'<div style="{style}font-size:{e["size"]}px;color:#{e["color"]};font-weight:{600 if e["bold"] else 400};text-align:{e["align"]};line-height:1.25;white-space:pre-wrap">{html.escape(e["text"])}</div>')
        elif e['kind']=='image':els.append(f'<img alt="Dayova App — aufgezeichneter Onboarding-Zustand" style="{style}object-fit:contain" src="data:image/png;base64,{base64.b64encode(Path(e["path"]).read_bytes()).decode()}">')
        else:els.append(f'<div style="{style}background:#{e["color"]};border-radius:{e.get("radius",999 if e["kind"]=="circle" else 0)}px"></div>')
    sections.append(f'<section data-document-role="page" data-label="{pg["label"]}" data-speaker-notes="{html.escape(pg["note"],quote=True)}" style="position:relative;width:1280px;height:720px;background:#{pg["bg"]};overflow:hidden">'+''.join(els)+'</section>')
(OUT/'Dayova-Canva.html').write_text('<!doctype html><html lang="de"><meta charset="utf-8"><title>Dayova — Vorlagen</title><style>'+fontcss+'*{box-sizing:border-box}body{margin:0;background:#ddd;font-family:Poppins}section{margin:24px auto}@media print{body{background:white}section{margin:0;break-after:page}@page{size:1280px 720px;margin:0}}</style>'+''.join(sections)+'</html>',encoding='utf-8', newline='\n')
# A matching vector reference PDF, also useful when fonts are not installed.
for name,file in [('Poppins','Poppins-Regular.ttf'),('PoppinsSemi','Poppins-SemiBold.ttf')]:pdfmetrics.registerFont(TTFont(name,str(ROOT/'assets/fonts'/file)))
pdf=canvas.Canvas(str(OUT/'Dayova-Reference.pdf'),pagesize=(W,H));pdf.setTitle('Dayova — template visual reference')
from reportlab.lib.colors import HexColor
for pg in pages:
    pdf.setFillColor(HexColor('#'+pg['bg']));pdf.rect(0,0,W,H,fill=1,stroke=0)
    for e in flattened(pg):
        x,y,w,h=[e[k] for k in ['x','y','w','h']]
        if e['kind']=='text':
            pdf.setFont('PoppinsSemi' if e['bold'] else 'Poppins',e['size']);pdf.setFillColor(HexColor('#'+e['color']))
            # Preserve explicit line breaks and wrap to the fixed text box.
            lines=[]
            for raw in e['text'].split('\n'):
                line=''
                for word in raw.split(' '):
                    test=(line+' '+word).strip()
                    if pdfmetrics.stringWidth(test,'PoppinsSemi' if e['bold'] else 'Poppins',e['size'])>w and line:lines.append(line);line=word
                    else:line=test
                lines.append(line)
            for i,line in enumerate(lines):
                if e.get('optical_center'):
                    top,bottom=ink_bounds(e.get('center_reference',line),e['bold'])
                    yy=H-y-h/2+(top+bottom)*e['size']/2
                else:yy=H-y-e['size']*.82-i*e['size']*1.25
                if e['align']=='center':pdf.drawCentredString(x+w/2,yy,line)
                elif e['align']=='right':pdf.drawRightString(x+w,yy,line)
                else:pdf.drawString(x,yy,line)
        elif e['kind']=='image':pdf.drawImage(e['path'],x,H-y-h,w,h,mask='auto')
        else:
            pdf.setFillColor(HexColor('#'+e['color']))
            if e['kind']=='circle':pdf.ellipse(x,H-y-h,x+w,H-y,fill=1,stroke=0)
            else:pdf.roundRect(x,H-y-h,w,h,min(e.get('radius',0),h/2,w/2),fill=1,stroke=0)
    pdf.showPage()
pdf.save()

# Workbook: bounded input area with robust blank / zero handling and portable formulas.
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
names=['Start','Dashboard','Lernschritte','Budget','Marke']
for name in names:
    ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.sheet_properties.tabColor=C['cyan'];ws.freeze_panes='A7' if name in ['Lernschritte','Budget'] else None;ws.sheet_view.zoomScale=90
    for row in ws.iter_rows(min_row=1,max_row=110,max_col=10):
        for cell in row:cell.fill=PatternFill('solid',fgColor=C['bg']);cell.font=Font(name='Poppins',size=11,color=C['ink']);cell.alignment=Alignment(vertical='center',wrap_text=True)
    for col in 'ABCDEFGHIJ':ws.column_dimensions[col].width=18
    ws.column_dimensions['A'].width=4;ws.column_dimensions['B'].width=30
    for row in range(1,111):ws.row_dimensions[row].height=28
    ws.merge_cells('B2:J3');ws['B2']='dayova.  /  '+name;ws['B2'].font=Font(name='Poppins',size=27,bold=True,color=C['ink']);ws.row_dimensions[2].height=35
    ws.merge_cells('B4:J4');ws['B4']='VORLAGE · Fiktive Beispieldaten · Stand 06.09.2026';ws['B4'].font=Font(name='Poppins',size=10,color=C['muted'])
    ws.print_options.horizontalCentered=True;ws.sheet_properties.outlinePr.summaryRight=False;ws.print_area='B2:J25';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
def band(ws,row,values):
    for i,v in enumerate(values,2):
        c=ws.cell(row,i,v);c.fill=PatternFill('solid',fgColor=C['ink']);c.font=Font(name='Poppins',bold=True,size=11,color=C['white'])
    ws.row_dimensions[row].height=38
def block(ws,range_,value,size=13,fill='white'):
    ws.merge_cells(range_)
    for row in ws[range_]:
        for cell in row:cell.fill=PatternFill('solid',fgColor=C[fill])
    c=ws[range_.split(':')[0]];c.value=value;c.font=Font(name='Poppins',size=size,color=C['ink'],bold=size>=20);c.alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
ws=wb['Start'];block(ws,'B6:J7','Einfach eintragen. Den Überblick behalten.',23,'ice')
for r,t,b in [(9,'01  Kopie erstellen','Datei → Kopie erstellen. Diese Vorlage als unverändertes Original behalten.'),(12,'02  Eingaben ersetzen','Lernschritte und Budget: hellblaue Zellen bearbeiten. Zeilen 7–106 sind vorbereitet. Beispielzeilen ersetzen oder Eingaben löschen.'),(15,'03  Überblick nutzen','Dashboard berechnet aus den Eingaben. Zeilen nur mit ID (Lernschritte) bzw. Position (Budget) zählen. Keine personenbezogenen Lernerdaten nötig.'),(18,'04  Vor dem Teilen prüfen','Beispieldaten, Datumsbereich, Formeln und Diagramm prüfen. Zahlen in Präsentationen manuell aktualisieren. Kein automatischer App- oder Linear-Abgleich.')]:
    block(ws,f'B{r}:J{r}',t,17);block(ws,f'B{r+1}:J{r+2}',b,12)
block(ws,'B23:J25','Zeilen 7–106 = 100 Datensätze. Darüber hinaus müssen Formeln, Filter und Validierungen erweitert werden. 0 % bedeutet: keine erfassten abgeschlossenen Lernschritte. Nur Blau = Eingabe; Grau/Weiß = berechnet oder Anleitung. Status wird zusätzlich als Text angezeigt.',12,'ice')
ws=wb['Lernschritte'];headers=['ID','Datum','Fach','Thema','Typ','Minuten','Status','Erledigt','Notiz'];band(ws,6,headers)
for col,width in [('B',15),('C',16),('D',19),('E',34),('F',20),('G',14),('H',22),('I',16),('J',33)]:ws.column_dimensions[col].width=width
sample=[('Mathematik','Lineare Funktionen','Theorie',30,'Erledigt'),('Biologie','Zellaufbau','Üben',25,'Erledigt'),('Englisch','Textverständnis','Praxis',20,'Geplant'),('Mathematik','Gleichungen','Üben',30,'In Arbeit'),('Deutsch','Argumentieren','Hausaufgabe',25,'Geplant'),('Biologie','Genetik','Theorie',30,'Erledigt')]
for r in range(7,107):
    for c in list(range(2,9))+[10]:ws.cell(r,c).fill=PatternFill('solid',fgColor=C['ice'])
    ws.cell(r,9,f'=IF(B{r}="","",IF(H{r}="Erledigt",1,0))').number_format='0%'
    for c in [7,9]:ws.cell(r,c).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    ws.cell(r,3).number_format='dd.mm.yyyy';ws.cell(r,7).number_format='0" Min."'
    if r<13:
        subj,topic,typ,mins,status=sample[r-7]
        for c,v in enumerate([f'LS-{r-6:03}',date(2026,9,7)+timedelta(days=r-7),subj,topic,typ,mins,status],2):ws.cell(r,c,v)
        ws.cell(r,10,'Fiktives Beispiel')
for rng,formula in [('F7:F106','"Theorie,Üben,Praxis,Hausaufgabe"'),('H7:H106','"Geplant,In Arbeit,Erledigt"')]:
    dv=DataValidation(type='list',formula1=formula,allow_blank=True);dv.errorTitle='Bitte einen Listenwert wählen';dv.error='Wähle einen Wert aus der Liste.';dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add(rng)
dv=DataValidation(type='whole',operator='between',formula1=0,formula2=1440,allow_blank=True);dv.showErrorMessage=True;dv.error='Minuten: ganze Zahl von 0 bis 1440.';ws.add_data_validation(dv);dv.add('G7:G106')
ws.auto_filter.ref='B6:J106';ws.print_title_rows='1:6';ws.print_area='B2:J16'
for label,color in [('Erledigt','palegreen'),('In Arbeit','paleorange')]:ws.conditional_formatting.add('H7:H106',FormulaRule(formula=[f'H7="{label}"'],fill=PatternFill('solid',fgColor=C[color])))
ws=wb['Budget'];band(ws,6,['Position','Kategorie','Menge','Einzelpreis','Plan','Ist','Rest','Status','Notiz'])
ws.column_dimensions['B'].width=29;ws.column_dimensions['C'].width=22;ws.column_dimensions['J'].width=34
for r in range(7,107):
    for c in [2,3,4,5,7,10]:ws.cell(r,c).fill=PatternFill('solid',fgColor=C['ice'])
    ws.cell(r,6,f'=IF(B{r}="","",IFERROR(D{r}*E{r},0))');ws.cell(r,8,f'=IF(B{r}="","",F{r}-G{r})');ws.cell(r,9,f'=IF(B{r}="","",IF(H{r}<0,"Über Plan","Im Plan"))')
    for c in range(5,9):ws.cell(r,c).number_format='#,##0.00" €";[Red]-#,##0.00" €"'
    if r<10:
        vals=[['Material','Durchführung',10,8,65],['Vorbereitung','Organisation',2,40,80],['Rückmeldung','Auswertung',1,50,60]][r-7]
        for c,v in zip([2,3,4,5,7],vals):ws.cell(r,c,v)
        ws.cell(r,10,'Fiktive Werte · keine Dayova Preise')
for rng in ['D7:E106','G7:G106']:
    dv=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1=0,allow_blank=True);dv.showErrorMessage=True;dv.error='Bitte eine Zahl ab 0 eingeben.';ws.add_data_validation(dv);dv.add(rng)
ws.auto_filter.ref='B6:J106';ws.conditional_formatting.add('I7:I106',FormulaRule(formula=['I7="Über Plan"'],fill=PatternFill('solid',fgColor=C['paleorange'])));ws.print_area='B2:J15'
ws=wb['Dashboard'];block(ws,'B6:J7','Der nächste Schritt beginnt mit Überblick.',23,'ice')
metrics=[('B9:D9','B10:D12','Lernschritte','=COUNTIF(Lernschritte!B7:B106,"<>")','0'),('E9:G9','E10:G12','Abgeschlossen','=IFERROR(SUM(Lernschritte!I7:I106)/COUNTIF(Lernschritte!B7:B106,"<>"),0)','0%'),('H9:J9','H10:J12','Geplante Minuten','=SUMIF(Lernschritte!B7:B106,"<>",Lernschritte!G7:G106)','0" Min."')]
for lr,vr,title,formula,fmt in metrics:block(ws,lr,title,12);block(ws,vr,formula,30);ws[vr.split(':')[0]].number_format=fmt
band(ws,15,['Status','Anzahl']);
for r,label in enumerate(['Geplant','In Arbeit','Erledigt'],16):ws.cell(r,2,label);ws.cell(r,3,f'=COUNTIFS(Lernschritte!B7:B106,"<>",Lernschritte!H7:H106,B{r})')
block(ws,'B21:D21','Budget · Plan',12);block(ws,'B22:D23','=SUM(Budget!F7:F106)',23);ws['B22'].number_format='#,##0.00" €"'
block(ws,'E21:G21','Budget · Ist',12);block(ws,'E22:G23','=SUMIF(Budget!B7:B106,"<>",Budget!G7:G106)',23);ws['E22'].number_format='#,##0.00" €"'
block(ws,'H21:J21','Budget · Rest',12);block(ws,'H22:J23','=B22-E22',23);ws['H22'].number_format='#,##0.00" €"'
ch=BarChart();ch.type='bar';ch.style=10;ch.title='Lernschritte nach Status';ch.add_data(Reference(ws,min_col=3,min_row=15,max_row=18),titles_from_data=True);ch.set_categories(Reference(ws,min_col=2,min_row=16,max_row=18));ch.height=5.4;ch.width=18;ch.legend=None;ch.x_axis.scaling.min=0;ch.series[0].graphicalProperties.solidFill=C['cyan'];ch.series[0].graphicalProperties.line.solidFill=C['cyan'];ws.add_chart(ch,'E15')
ws=wb['Marke'];band(ws,6,['Rolle','HEX','Verwendung'])
for r,(label,key,use) in enumerate([('Primär','cyan','Akzente; dunkler Text'),('Text','ink','Primärtext'),('Hintergrund','bg','Arbeitsfläche'),('Fläche','white','Karten'),('Eingabe','ice','Bearbeitbare Zellen'),('Sekundärtext','muted','Hinweise'),('Theorie','purple','Kategorie, nicht Warnstatus')],7):
    ws.cell(r,2,label);ws.cell(r,3,'#'+C[key]);ws.cell(r,4,use);ws.cell(r,3).fill=PatternFill('solid',fgColor=C[key]);ws.cell(r,3).font=Font(name='Poppins',color=C['white'] if key in ['ink','muted','purple'] else C['ink'])
ws.column_dimensions['D'].width=39
block(ws,'B16:J18','Poppins Regular für Fließtext, Poppins SemiBold für Überschriften. In Sheets wird Hervorhebung über Fett dargestellt. #00BAFF immer mit #1A1A1A kombinieren. Kategorien nie als einzige Statusinformation nutzen.',13,'ice')
block(ws,'B20:J23','Quellen: dayova.com (06.09.2026), App-Tokens und Design-System-Kontext im Repository. Figma ist Referenz, kein Freigabegate. Die Entscheidungsdokumentation und Nutzungshinweise liegen im Kit. Eingabeflächen sind formatiert, nicht gegen Änderungen gesperrt.',12)
wb.active=0;wb.save(OUT/'Dayova-Workbook.xlsx')
(OUT/'brand-tokens.json').write_text(json.dumps({'version':'1.0','date':'2026-09-06','colors':C,'font':'Poppins','presentation':{'width':1280,'height':720,'margin':56,'bodyPx':20,'titlePx':44},'scope':'Communication templates; does not change app tokens'},indent=2),encoding='utf-8', newline='\n')
import pymupdf
doc=pymupdf.open(OUT/'Dayova-Reference.pdf');thumbs=[]
for i,pg in enumerate(doc):
    pix=pg.get_pixmap(matrix=pymupdf.Matrix(.4,.4));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);thumbs.append(im)
sheet=Image.new('RGB',(512*4,288*5),(220,220,220))
for i,im in enumerate(thumbs):sheet.paste(im,((i%4)*512,(i//4)*288))
sheet.save(OUT/'Dayova-Overview.png')
print(json.dumps({'slides':len(pages),'sheets':wb.sheetnames,'output':str(OUT)},indent=2))
